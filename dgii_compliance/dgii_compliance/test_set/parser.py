"""Parser for the DGII certification test workbook.

The workbook (``133424959-...xlsx``) holds one case per row with the expected
e-CF values laid out as flat columns.  Repeating groups use bracket suffixes:

* ``Col[n]``     -> one value per line *or* per header table row
* ``Col[n][m]``  -> nested value (line ``n``, sub-index ``m``)

The parser only normalizes; deciding what is a header table vs. a line field is
done in :mod:`mapping`.
"""

import datetime
import re

import openpyxl
from frappe.utils import cint, flt

EMPTY_VALUES = (None, "", "#e")
SINGLE_RE = re.compile(r"^(?P<base>.+)\[(?P<idx>\d+)\]$")
DOUBLE_RE = re.compile(r"^(?P<base>.+)\[(?P<line>\d+)\]\[(?P<sub>\d+)\]$")

DEFAULT_SHEET = "ECF"


def _clean(value):
	if value is None:
		return None
	if isinstance(value, str):
		value = value.strip()
		if value in EMPTY_VALUES:
			return None
	return value


def to_date(value):
	"""Parses a workbook date (``DD-MM-YYYY`` string or date/datetime)."""
	value = _clean(value)
	if value is None:
		return None
	if isinstance(value, datetime.datetime):
		return value.date()
	if isinstance(value, datetime.date):
		return value
	for fmt in ("%d-%m-%Y", "%Y-%m-%d"):
		try:
			return datetime.datetime.strptime(str(value), fmt).date()
		except ValueError:
			continue
	return None


def to_float(value):
	value = _clean(value)
	return None if value is None else flt(value)


def to_int(value):
	value = _clean(value)
	return None if value is None else cint(value)


def _read_header(worksheet):
	header = {}
	for col in range(1, worksheet.max_column + 1):
		name = worksheet.cell(row=1, column=col).value
		if name is None:
			continue
		# A couple of columns carry a stray trailing space.
		header[col] = str(name).strip()
	return header


def parse_case(raw):
	"""Parses a single worksheet row into a normalized case dict."""
	row, header = raw
	case = {"header": {}, "indexed": {}, "sub": {}}
	for col, name in header.items():
		value = _clean(row[col - 1])
		if value is None:
			continue
		double = DOUBLE_RE.match(name)
		if double:
			line = int(double.group("line"))
			sub = int(double.group("sub"))
			case["sub"].setdefault(double.group("base"), {}).setdefault(line, {})[sub] = value
			continue
		single = SINGLE_RE.match(name)
		if single:
			case["indexed"].setdefault(single.group("base"), {})[int(single.group("idx"))] = value
			continue
		case["header"][name] = value

	case["case_id"] = case["header"].get("CasoPrueba")
	tipo = case["header"].get("TipoeCF")
	case["tipo"] = f"E{tipo}" if tipo is not None else None
	case["encf"] = case["header"].get("ENCF")
	return case


def parse_workbook(path, sheet=DEFAULT_SHEET):
	"""Returns the list of normalized cases from ``path``/``sheet``."""
	worksheet = openpyxl.load_workbook(path, data_only=True, read_only=True)[sheet]
	header = _read_header(worksheet)
	cases = []
	for row in worksheet.iter_rows(min_row=2, values_only=True):
		if not any(cell not in EMPTY_VALUES for cell in row):
			continue
		case = parse_case((row, header))
		if case["encf"]:
			cases.append(case)
	return cases


def dedupe(cases):
	"""Keeps the first occurrence of each e-NCF (ECF sheet is authoritative)."""
	seen = {}
	for case in cases:
		seen.setdefault(case["encf"], case)
	return list(seen.values())


def needs_rfce(case):
	"""True for E32 consumption invoices below RD$250,000 (Resumen path)."""
	return case["tipo"] == "E32" and (to_float(case["header"].get("MontoTotal")) or 0) < 250000


def summarize(cases):
	by_type = {}
	rfce = 0
	for case in cases:
		by_type[case["tipo"]] = by_type.get(case["tipo"], 0) + 1
		if needs_rfce(case):
			rfce += 1
	return {
		"total": len(cases),
		"by_type": dict(sorted(by_type.items())),
		"rfce": rfce,
		"ecf": len(cases) - rfce,
	}
